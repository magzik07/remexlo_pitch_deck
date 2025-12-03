#!/usr/bin/env python3
"""
Script to generate a PPTX pitch deck for Remexlo - 
A marketplace platform connecting craftsmen with customers.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


def create_title_slide(prs):
    """Create the title slide."""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Add background shape
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(26, 35, 126)  # Dark blue
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "REMEXLO"
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Tagline
    tagline_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9), Inches(1))
    tf = tagline_box.text_frame
    p = tf.paragraphs[0]
    p.text = "The Marketplace for Craftsmen"
    p.font.size = Pt(32)
    p.font.color.rgb = RGBColor(255, 193, 7)  # Amber/Gold
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(9), Inches(0.5))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Connecting Skilled Artisans with Customers Worldwide"
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(200, 200, 200)
    p.alignment = PP_ALIGN.CENTER


def create_problem_slide(prs):
    """Create the problem statement slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "The Problem"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 35, 126)
    
    # Problems list
    problems = [
        "🔧 Skilled craftsmen struggle to find customers and market their services",
        "🏠 Homeowners find it difficult to locate reliable, quality craftsmen",
        "💰 Lack of transparent pricing leads to distrust on both sides",
        "📱 No dedicated digital platform for skilled trades",
        "⭐ No standardized way to verify quality and build reputation"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, problem in enumerate(problems):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = problem
        p.font.size = Pt(24)
        p.space_after = Pt(20)


def create_solution_slide(prs):
    """Create the solution slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Our Solution: Remexlo"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 35, 126)
    
    # Solution points
    solutions = [
        "✅ Digital marketplace connecting craftsmen with customers",
        "✅ Verified profiles with portfolio showcases",
        "✅ Transparent pricing and instant quotes",
        "✅ Rating and review system for quality assurance",
        "✅ Secure payment processing with escrow",
        "✅ Mobile-first approach for easy access"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, solution in enumerate(solutions):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = solution
        p.font.size = Pt(24)
        p.space_after = Pt(18)


def create_market_slide(prs):
    """Create the market opportunity slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Market Opportunity"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 35, 126)
    
    # Market stats boxes
    stats = [
        ("$550B", "Global Home Services Market"),
        ("12%", "Annual Market Growth"),
        ("78M", "Active Craftsmen Worldwide"),
        ("68%", "Customers Prefer Online Booking")
    ]
    
    left_positions = [0.5, 5]
    top_positions = [1.5, 3.3]
    
    for idx, (value, label) in enumerate(stats):
        left = left_positions[idx % 2]
        top = top_positions[idx // 2]
        
        # Value
        value_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(4), Inches(0.8))
        tf = value_box.text_frame
        p = tf.paragraphs[0]
        p.text = value
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.color.rgb = RGBColor(26, 35, 126)
        
        # Label
        label_box = slide.shapes.add_textbox(Inches(left), Inches(top + 0.8), Inches(4), Inches(0.5))
        tf = label_box.text_frame
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(100, 100, 100)


def create_business_model_slide(prs):
    """Create the business model slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Business Model"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 35, 126)
    
    # Revenue streams
    revenues = [
        ("Transaction Fees", "8-12% commission on completed jobs"),
        ("Premium Listings", "Featured placement for craftsmen - $29/month"),
        ("Subscription Plans", "Pro tools for businesses - $49-199/month"),
        ("Lead Generation", "Verified customer leads - pay per lead")
    ]
    
    y_pos = 1.5
    for title, desc in revenues:
        # Title
        t_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(4), Inches(0.5))
        tf = t_box.text_frame
        p = tf.paragraphs[0]
        p.text = "💵 " + title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(26, 35, 126)
        
        # Description
        d_box = slide.shapes.add_textbox(Inches(0.8), Inches(y_pos + 0.4), Inches(8), Inches(0.4))
        tf = d_box.text_frame
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(100, 100, 100)
        
        y_pos += 1.0


def create_traction_slide(prs):
    """Create the traction/roadmap slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Traction & Roadmap"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 35, 126)
    
    # Traction
    traction_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(4), Inches(2.5))
    tf = traction_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Current Traction"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 193, 7)
    
    traction_items = ["5,000+ registered craftsmen", "15,000+ active users", 
                      "$250K GMV to date", "4.8★ average rating"]
    for item in traction_items:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(18)
        p.space_after = Pt(8)
    
    # Roadmap
    roadmap_box = slide.shapes.add_textbox(Inches(5), Inches(1.3), Inches(4.5), Inches(2.5))
    tf = roadmap_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Roadmap"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 193, 7)
    
    roadmap_items = [
        "Q1: Mobile app launch",
        "Q2: Expand to 5 new cities",
        "Q3: AI-powered matching",
        "Q4: International expansion"
    ]
    for item in roadmap_items:
        p = tf.add_paragraph()
        p.text = "→ " + item
        p.font.size = Pt(18)
        p.space_after = Pt(8)


def create_team_slide(prs):
    """Create the team slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Our Team"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 35, 126)
    
    # Team members
    team = [
        ("👤 CEO & Co-Founder", "10+ years in marketplace startups\nPreviously at Upwork"),
        ("👤 CTO & Co-Founder", "Former tech lead at Amazon\nBuilt platforms serving 10M+ users"),
        ("👤 Head of Operations", "Operations expert from TaskRabbit\n5+ years scaling home services"),
        ("👤 Head of Growth", "Growth marketing from Thumbtack\nDriven 10x user acquisition")
    ]
    
    positions = [(0.5, 1.5), (5, 1.5), (0.5, 3.3), (5, 3.3)]
    
    for idx, ((role, desc), (left, top)) in enumerate(zip(team, positions)):
        # Role
        role_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(4), Inches(0.5))
        tf = role_box.text_frame
        p = tf.paragraphs[0]
        p.text = role
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = RGBColor(26, 35, 126)
        
        # Description
        desc_box = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.5), Inches(4), Inches(1))
        tf = desc_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(100, 100, 100)


def create_ask_slide(prs):
    """Create the investment ask slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(26, 35, 126)
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "The Ask"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Amount
    amount_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(9), Inches(1))
    tf = amount_box.text_frame
    p = tf.paragraphs[0]
    p.text = "$2.5M Seed Round"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 193, 7)
    p.alignment = PP_ALIGN.CENTER
    
    # Use of funds
    funds_box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(2.5))
    tf = funds_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Use of Funds:"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    funds = ["40% - Product Development & Engineering",
             "30% - Sales & Marketing",
             "20% - Operations & Customer Success",
             "10% - General & Administrative"]
    
    for fund in funds:
        p = tf.add_paragraph()
        p.text = "• " + fund
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(200, 200, 200)
        p.space_after = Pt(8)


def create_contact_slide(prs):
    """Create the closing/contact slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(26, 35, 126)
    background.line.fill.background()
    
    # Thank you
    thank_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(1))
    tf = thank_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Thank You"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Logo
    logo_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(9), Inches(0.8))
    tf = logo_box.text_frame
    p = tf.paragraphs[0]
    p.text = "REMEXLO"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 193, 7)
    p.alignment = PP_ALIGN.CENTER
    
    # Contact info
    contact_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(9), Inches(1.5))
    tf = contact_box.text_frame
    
    p = tf.paragraphs[0]
    p.text = "📧 invest@remexlo.com"
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(200, 200, 200)
    p.alignment = PP_ALIGN.CENTER
    
    p = tf.add_paragraph()
    p.text = "🌐 www.remexlo.com"
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(200, 200, 200)
    p.alignment = PP_ALIGN.CENTER


def main():
    """Generate the pitch deck."""
    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9 aspect ratio
    
    # Create all slides
    create_title_slide(prs)
    create_problem_slide(prs)
    create_solution_slide(prs)
    create_market_slide(prs)
    create_business_model_slide(prs)
    create_traction_slide(prs)
    create_team_slide(prs)
    create_ask_slide(prs)
    create_contact_slide(prs)
    
    # Save presentation
    output_path = "Remexlo_Pitch_Deck.pptx"
    prs.save(output_path)
    print(f"Pitch deck successfully created: {output_path}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
